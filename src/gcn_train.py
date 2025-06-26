import random
import numpy as np
import os
import torch
import torch.nn as nn
import torch.nn.functional as F # For get_loss
import pickle
import openpyxl # 新增导入
from openpyxl.utils import get_column_letter # 新增导入，用于调整列宽
import math # 新增导入，用于计算分割点

from evaluator import evaluate # Assuming evaluator.py is in the same directory or accessible
from gcn_model import GCNStockPredictor # Import GCN model

# Loss function (copied here for self-containment, or could be in a utils.py)
def get_loss(prediction, ground_truth, base_price, mask, batch_size_stock_num, alpha):
    device = prediction.device
    all_one = torch.ones(batch_size_stock_num, 1, dtype=torch.float32).to(device)
    
    # 保守的形状检查和修复：只在检测到可能的广播问题时进行修复
    original_pred_shape = prediction.shape
    original_base_shape = base_price.shape
    
    # 检查是否可能出现广播问题
    if prediction.ndim == 1 and len(base_price.shape) == 2 and base_price.shape[1] == 1:
        # 这种情况下会发生 (N,) vs (N,1) 的广播，转换为 (N,1) vs (N,1)
        prediction = prediction.unsqueeze(1)
        print(f"信息：检测到潜在的广播问题，已修复预测张量形状: {original_pred_shape} -> {prediction.shape}")
    elif prediction.ndim == 2 and prediction.shape != base_price.shape:
        # 其他可能的形状不匹配情况
        if prediction.shape[0] == base_price.shape[0] and prediction.shape[1] != 1:
            print(f"警告：预测张量形状异常: {original_pred_shape}，base_price形状: {original_base_shape}")
            prediction = prediction.squeeze().unsqueeze(1)
            print(f"已尝试修复为: {prediction.shape}")
    
    return_ratio = torch.div(torch.sub(prediction, base_price), base_price + 1e-9) # Added epsilon to base_price
    reg_loss = F.mse_loss(return_ratio * mask, ground_truth * mask)
    pre_pw_dif = torch.sub(
        return_ratio @ all_one.t(),
        all_one @ return_ratio.t()
    )
    gt_pw_dif = torch.sub(
        all_one @ ground_truth.t(),
        ground_truth @ all_one.t()
    )
    mask_pw = mask @ mask.t()
    rank_loss = torch.mean(
        F.relu(pre_pw_dif * gt_pw_dif * mask_pw)
    )
    loss = reg_loss + alpha * rank_loss
    return loss, reg_loss, rank_loss, return_ratio


np.random.seed(123456789)
torch.manual_seed(12345678)
device = torch.device("cuda" if torch.cuda.is_available() else 'cpu')
print(f"Using device: {device}")

# Define paths and basic parameters (mirroring previous setup)
data_path = '../dataset' # Relative to src directory where this script is
market_name = 'A_SHARE' 
stock_num = 1026 # For NASDAQ, will be dynamically updated after data loading
lookback_length = 16
epochs = 100 # Consistent with previous settings
# valid_index and test_index will be calculated dynamically after data loading
fea_num = 5 # Input feature dimension for GCN (from the last time step)
steps = 1 # For ground truth calculation offset
learning_rate = 0.001
alpha_loss = 0.1 # Weight for rank loss in get_loss
accumulation_steps = 16 # Gradient accumulation steps

# GCN specific hyperparameters
num_gcn_layers = 2
gcn_hidden_dims = [32, 16] # Example: 2 layers, 32 then 16 features
gcn_dropout_rate = 0.1
gcn_activation_fn_class = nn.ReLU
gcn_use_tanh_output = True # Recommended to keep True to avoid NaN IC/RIC

# Load data
dataset_path = os.path.join(data_path, market_name)
print(f"Loading data from: {dataset_path}")

if market_name == "SP500":
    # SP500数据集的特殊加载和预处理逻辑
    data = np.load('../dataset/SP500/SP500.npy')
    print(f"原始SP500数据形状: {data.shape} (股票数: {data.shape[0]}, 天数: {data.shape[1]}, 特征数: {data.shape[2]})")
    
    # 使用全部5年数据，不再进行切片
    print(f"使用全部5年数据 (2020-2024)，数据形状: {data.shape}")
    
    price_data = data[:, :, -1] # 价格数据
    mask_data = np.ones((data.shape[0], data.shape[1])) # 掩码数据
    eod_data = data # EOD数据
    gt_data = np.zeros((data.shape[0], data.shape[1])) # 真实收益率数据初始化
    # 计算真实收益率
    for ticket in range(0, data.shape[0]):
        for row in range(1, data.shape[1]):
            gt_data[ticket][row] = (data[ticket][row][-1] - data[ticket][row - steps][-1]) / \
                                   (data[ticket][row - steps][-1] + 1e-8) # 添加小值避免除零
else:
    # 其他市场 (如NASDAQ, NYSE) 的数据加载逻辑
    try:
        with open(os.path.join(dataset_path, "eod_data.pkl"), "rb") as f:
            eod_data = pickle.load(f)
        with open(os.path.join(dataset_path, "mask_data.pkl"), "rb") as f:
            mask_data = pickle.load(f)
        with open(os.path.join(dataset_path, "gt_data.pkl"), "rb") as f:
            gt_data = pickle.load(f)
        with open(os.path.join(dataset_path, "price_data.pkl"), "rb") as f:
            price_data = pickle.load(f)
    except FileNotFoundError as e:
        print(f"Error loading data files: {e}")
        print(f"Please ensure {dataset_path}/ contains eod_data.pkl, mask_data.pkl, gt_data.pkl, price_data.pkl")
        exit()

# Verify stock_num if data is loaded (e.g., from eod_data shape)
if eod_data.shape[0] != stock_num:
    print(f"Warning: stock_num ({stock_num}) does not match eod_data.shape[0] ({eod_data.shape[0]}). Adjusting stock_num.")
    stock_num = eod_data.shape[0]

trade_dates = mask_data.shape[1]

# --- 通用设置和动态分割 ---
print(f"市场为 {market_name}，实际加载股票数量更新为: {stock_num}")
print(f"总交易天数: {trade_dates}")

# 动态计算数据集分割点
train_ratio = 0.60
valid_ratio = 0.20
# 测试集比例由 (1 - train_ratio - valid_ratio) 隐式确定

# 确保最小数据集大小
min_valid_days = max(50, lookback_length + steps + 10)  # 验证集最少50天或lookback+steps+10天
min_test_days = max(50, lookback_length + steps + 10)   # 测试集最少50天或lookback+steps+10天
min_train_days = max(100, lookback_length + steps + 20) # 训练集最少100天或lookback+steps+20天

# 计算分割点
valid_index = max(min_train_days, math.floor(trade_dates * train_ratio))
test_index = max(valid_index + min_valid_days, math.floor(trade_dates * (train_ratio + valid_ratio)))

# 确保测试集也有足够的数据
if trade_dates - test_index < min_test_days:
    # 如果测试集太小，向前调整
    test_index = trade_dates - min_test_days
    if test_index <= valid_index:
        # 如果还是不够，重新分配
        available_days = trade_dates - min_train_days
        if available_days >= min_valid_days + min_test_days:
            valid_index = min_train_days
            test_index = trade_dates - min_test_days
        else:
            print(f"错误：数据集太小({trade_dates}天)，无法进行有效的训练/验证/测试分割")
            print(f"最少需要: {min_train_days + min_valid_days + min_test_days}天")
            exit(1)

print(f"根据 {train_ratio:.0%}/{valid_ratio:.0%}/_ 分割比例 (已调整为满足最小数据集要求):")
print(f"训练集天数 (0-indexed, up to valid_index-1): {valid_index} (实际比例: {valid_index/trade_dates:.1%})")
print(f"验证集天数 (valid_index to test_index-1): {test_index - valid_index} (实际比例: {(test_index-valid_index)/trade_dates:.1%})")
print(f"测试集天数 (test_index to end): {trade_dates - test_index} (实际比例: {(trade_dates-test_index)/trade_dates:.1%})")
print(f"验证集起始索引 (valid_index): {valid_index}")
print(f"测试集起始索引 (test_index): {test_index}")

# 最终安全检查
if valid_index >= test_index or test_index >= trade_dates:
    print(f"错误：分割索引无效 - valid_index:{valid_index}, test_index:{test_index}, trade_dates:{trade_dates}")
    exit(1)
    
if test_index - valid_index <= 0:
    print(f"错误：验证集大小为0 - valid_index:{valid_index}, test_index:{test_index}")
    exit(1)
    
if trade_dates - test_index <= 0:
    print(f"错误：测试集大小为0 - test_index:{test_index}, trade_dates:{trade_dates}")
    exit(1)

print("✅ 数据集分割检查通过")
# --- 结束通用逻辑 ---

# Initialize GCNStockPredictor model
model = GCNStockPredictor(
    num_stocks=stock_num,
    input_feature_dim=fea_num,
    num_gcn_layers=num_gcn_layers,
    gcn_hidden_dims=gcn_hidden_dims,
    output_dim=1, # Predicting a single value (e.g., stock return)
    dropout_rate=gcn_dropout_rate,
    activation_fn_class=gcn_activation_fn_class,
    use_tanh_output=gcn_use_tanh_output
).to(device)

optimizer = torch.optim.Adam(model.parameters(), lr=learning_rate)
best_valid_loss = float('inf')
best_valid_perf = None
best_test_perf = None
best_epoch_num = None # Initialize best_epoch_num

# Batching offsets for training (up to valid_index)
batch_offsets = np.arange(start=0, stop=valid_index - lookback_length - steps + 1, dtype=int)
if len(batch_offsets) == 0:
    print("Error: Not enough data for training based on valid_index, lookback_length, and steps.")
    print(f"valid_index={valid_index}, lookback_length={lookback_length}, steps={steps}")
    exit()


def get_batch(offset): # offset is the start of the lookback window
    seq_len = lookback_length
    # EOD data: (stock_num, lookback_length, fea_num)
    eod_batch = eod_data[:, offset:offset + seq_len, :]
    
    # Mask data: (stock_num, lookback_length + steps)
    # The mask for a sample should ensure all days in lookback AND the day for GT are valid.
    mask_raw = mask_data[:, offset: offset + seq_len + steps]
    # If any day in this window is masked for a stock, that stock is masked for this sample.
    current_mask_batch = np.min(mask_raw, axis=1)
    
    # Price data for loss calculation (price at t-1 relative to prediction at t)
    # This is the price at the END of the lookback window, used as base for return_ratio.
    current_price_batch = price_data[:, offset + seq_len - 1]
    
    # Ground truth is for `steps` ahead of the end of the lookback window.
    current_gt_batch = gt_data[:, offset + seq_len + steps - 1]
    
    return (
        eod_batch,
        np.expand_dims(current_mask_batch, axis=1),
        np.expand_dims(current_price_batch, axis=1),
        np.expand_dims(current_gt_batch, axis=1)
    )

def validate(start_eval_index, end_eval_index, phase_name="Valid"):
    model.eval()
    with torch.no_grad():
        cur_pred = np.zeros([stock_num, end_eval_index - start_eval_index], dtype=float)
        cur_gt = np.zeros([stock_num, end_eval_index - start_eval_index], dtype=float)
        cur_mask = np.zeros([stock_num, end_eval_index - start_eval_index], dtype=float)
        
        total_loss, total_reg_loss, total_rank_loss = 0., 0., 0.
        
        # Offsets for validation/testing period
        # Prediction for day `t_predict` uses data up to `t_predict - steps` (end of lookback)
        # So, offset = `t_predict - steps - lookback_length`
        eval_offsets = np.arange(start_eval_index - lookback_length - steps + 1, 
                                 end_eval_index - lookback_length - steps + 1, # +1 because loop is exclusive for daily eval
                                 dtype=int)
        if len(eval_offsets) == 0:
            print(f"Warning: Not enough data for {phase_name} evaluation. Skipping.")
            return 0,0,0,None

        for day_idx, current_offset in enumerate(eval_offsets):
            data_b, mask_b, price_b, gt_b = map(
                lambda x: torch.from_numpy(x).float().to(device),
                get_batch(current_offset)
            )
            
            prediction_output = model(data_b) # Model output: (stock_num, 1)
            
            loss_val, reg_loss_val, rank_loss_val, rr_val = get_loss(
                prediction_output, gt_b, price_b, mask_b, stock_num, alpha_loss
            )
            total_loss += loss_val.item()
            total_reg_loss += reg_loss_val.item()
            total_rank_loss += rank_loss_val.item()
            
            # Store predictions (return_ratio) for evaluation
            cur_pred[:, day_idx] = rr_val[:, 0].cpu().numpy()
            cur_gt[:, day_idx] = gt_b[:, 0].cpu().numpy()
            cur_mask[:, day_idx] = mask_b[:, 0].cpu().numpy()

        num_days_evaluated = len(eval_offsets)
        avg_loss = total_loss / num_days_evaluated
        avg_reg_loss = total_reg_loss / num_days_evaluated
        avg_rank_loss = total_rank_loss / num_days_evaluated
        
        perf = evaluate(cur_pred, cur_gt, cur_mask)
    model.train()
    return avg_loss, avg_reg_loss, avg_rank_loss, perf

# Training loop
print("Starting training...")
for epoch in range(epochs):
    print(f"Epoch {epoch + 1}/{epochs} ##########################################################")
    np.random.shuffle(batch_offsets)
    
    tra_loss, tra_reg_loss, tra_rank_loss = 0., 0., 0.
    optimizer.zero_grad()
    
    num_train_batches = len(batch_offsets)
    for batch_idx, current_offset in enumerate(batch_offsets):
        data_b, mask_b, price_b, gt_b = map(
            lambda x: torch.from_numpy(x).float().to(device),
            get_batch(current_offset)
        )
        
        model.train()
        prediction_output = model(data_b)
        
        cur_loss, cur_reg_loss, cur_rank_loss, _ = get_loss(
            prediction_output, gt_b, price_b, mask_b, stock_num, alpha_loss
        )
        
        loss_to_backward = cur_loss / accumulation_steps
        loss_to_backward.backward()
        
        if (batch_idx + 1) % accumulation_steps == 0 or (batch_idx + 1) == num_train_batches:
            optimizer.step()
            optimizer.zero_grad()
            
        tra_loss += cur_loss.item()
        tra_reg_loss += cur_reg_loss.item()
        tra_rank_loss += cur_rank_loss.item()
        
    tra_loss /= num_train_batches
    tra_reg_loss /= num_train_batches
    tra_rank_loss /= num_train_batches
    print(f'Train: loss:{tra_loss:.4e} = {tra_reg_loss:.4e} + alpha*{tra_rank_loss:.4e}')

    val_loss, val_reg_loss, val_rank_loss, val_perf = validate(valid_index, test_index, phase_name="Valid")
    print(f'Valid: loss:{val_loss:.4e} = {val_reg_loss:.4e} + alpha*{val_rank_loss:.4e}')
    if val_perf:
        print(f'  Perf: mse:{val_perf["mse"]:.4e}, IC:{val_perf["IC"]:.4f}, RIC:{val_perf["RIC"]:.4f}, prec@10:{val_perf["prec_10"]:.4f}, SR:{val_perf["sharpe5"]:.4f}')

    test_loss, test_reg_loss, test_rank_loss, test_perf = validate(test_index, trade_dates, phase_name="Test")
    print(f'Test : loss:{test_loss:.4e} = {test_reg_loss:.4e} + alpha*{test_rank_loss:.4e}')
    if test_perf:
        print(f'  Perf: mse:{test_perf["mse"]:.4e}, IC:{test_perf["IC"]:.4f}, RIC:{test_perf["RIC"]:.4f}, prec@10:{test_perf["prec_10"]:.4f}, SR:{test_perf["sharpe5"]:.4f}\n')

    if val_loss < best_valid_loss:
        best_valid_loss = val_loss
        best_valid_perf = val_perf
        best_test_perf = test_perf
        best_epoch_num = epoch + 1 # Record the epoch number
        # torch.save(model.state_dict(), best_model_path) # Line removed
        print(f"New best validation loss found at epoch {best_epoch_num}") # Updated print message

print("Training finished.")
if best_valid_perf and best_test_perf and best_epoch_num is not None: # Check best_epoch_num
    print(f"\nTraining finished. Best validation performance was at epoch {best_epoch_num}.") # Updated print
    print("Best Validation Performance (based on lowest validation loss):")
    print(f'mse:{best_valid_perf["mse"]:.4e}, IC:{best_valid_perf["IC"]:.4f}, RIC:{best_valid_perf["RIC"]:.4f}, prec@10:{best_valid_perf["prec_10"]:.4f}, SR:{best_valid_perf["sharpe5"]:.4f}')
    print("Corresponding Test Performance:")
    print(f'mse:{best_test_perf["mse"]:.4e}, IC:{best_test_perf["IC"]:.4f}, RIC:{best_test_perf["RIC"]:.4f}, prec@10:{best_test_perf["prec_10"]:.4f}, SR:{best_test_perf["sharpe5"]:.4f}')

    # --- 开始写入Excel的逻辑 ---
    model_name_for_excel = "GCN"
    excel_file_path = '../training_results.xlsx'
    header = ["Model", "Dataset", "Best Epoch",
              "Valid MSE", "Valid IC", "Valid RIC", "Valid Prec@10", "Valid SR",
              "Test MSE", "Test IC", "Test RIC", "Test Prec@10", "Test SR"]
    
    data_row = [model_name_for_excel, market_name, best_epoch_num,
                best_valid_perf.get('mse', float('nan')), 
                best_valid_perf.get('IC', float('nan')),
                best_valid_perf.get('RIC', float('nan')),
                best_valid_perf.get('prec_10', float('nan')),
                best_valid_perf.get('sharpe5', float('nan')),
                best_test_perf.get('mse', float('nan')),
                best_test_perf.get('IC', float('nan')),
                best_test_perf.get('RIC', float('nan')),
                best_test_perf.get('prec_10', float('nan')),
                best_test_perf.get('sharpe5', float('nan'))
               ]

    try:
        if not os.path.exists(excel_file_path):
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.title = "Training Results"
            sheet.append(header)
            # 自动调整列宽
            for col_idx, column_cells in enumerate(sheet.columns):
                length = max(len(str(cell.value)) for cell in column_cells)
                sheet.column_dimensions[get_column_letter(col_idx + 1)].width = length + 2
        else:
            workbook = openpyxl.load_workbook(excel_file_path)
            sheet = workbook.active
        
        sheet.append(data_row)
        workbook.save(excel_file_path)
        print(f"Results appended to {excel_file_path}")
    except Exception as e:
        print(f"Error writing to Excel: {e}")
    # --- 结束写入Excel的逻辑 ---

else:
    print("No best performance recorded. Check validation logic or training duration.") 