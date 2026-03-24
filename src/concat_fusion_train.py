#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
简单拼接多模态融合股票预测训练脚本（使用真实数据）

特点：
1. 融合两种模态：价格/技术特征 + 新闻情绪特征
2. 使用简单拼接方式融合两种模态（对比实验基准）
3. 使用与 gated_multimodal_train.py 相同的真实数据源与划分方式
4. 将最优运行结果写入 xlsx 文件
"""

import os
import math
import random
import pickle
import json
from typing import List, Dict, Tuple, Optional

import numpy as np
import pandas as pd
import torch
import torch.nn as nn
import torch.nn.functional as F
from torch.utils.data import Dataset, DataLoader

import openpyxl
from openpyxl.utils import get_column_letter

from evaluator import evaluate


# ===================== 全局配置 =====================

MARKET_NAME = "SP500"
LOOKBACK_LENGTH = 32          # 价格数据序列长度
PRED_STEPS = 1                # 预测步长
DATA_PATH = "../dataset"

BATCH_SIZE = 64             # 增加 batch size
EPOCHS = 80
LEARNING_RATE = 1e-4        # 回退到上一版稳定配置
WEIGHT_DECAY = 5e-4         # 回退到上一版稳定配置

# 排名损失权重（与第一部分 StockPredict 保持一致）
ALPHA = 0.5                 # 平衡回归损失和排名损失

# 学习率调度器参数
WARMUP_EPOCHS = 10
SCHEDULER_T_0 = 30
PATIENCE = 30               # 增加早停耐心

# 价格/技术特征编码器参数
PRICE_ENCODER_HIDDEN = 64
PRICE_ENCODER_LAYERS = 2
PRICE_DROPOUT = 0.2         # 增加 dropout

# 情绪特征编码器参数（回退到上一版）
SENTIMENT_ENCODER_HIDDEN = 32
SENTIMENT_ENCODER_LAYERS = 2
SENTIMENT_DROPOUT = 0.2

# 拼接融合参数（回退到上一版）
CONCAT_HIDDEN = 64

RANDOM_SEED = 1234


# ===================== 排名损失函数 =====================

def compute_rank_loss(prediction: torch.Tensor, ground_truth: torch.Tensor, batch_size: int) -> torch.Tensor:
    """
    计算成对排名损失 (Pairwise Ranking Loss)
    与 StockPredict 模型保持一致
    """
    device = prediction.device
    all_one = torch.ones(batch_size, 1, dtype=torch.float32).to(device)
    
    # 预测的成对差异
    pre_pw_dif = prediction @ all_one.t() - all_one @ prediction.t()
    # 真实值的成对差异
    gt_pw_dif = ground_truth @ all_one.t() - all_one @ ground_truth.t()
    
    # 排名损失：鼓励预测差异与真实差异方向一致
    rank_loss = torch.mean(F.relu(-pre_pw_dif * gt_pw_dif))
    return rank_loss


def compute_combined_loss(prediction: torch.Tensor, ground_truth: torch.Tensor, 
                          criterion: nn.MSELoss, alpha: float = 0.5) -> Tuple[torch.Tensor, torch.Tensor, torch.Tensor]:
    """
    计算组合损失：MSE + alpha * 排名损失
    """
    batch_size = prediction.size(0)
    
    # MSE 回归损失
    reg_loss = criterion(prediction, ground_truth)
    
    # 排名损失
    rank_loss = compute_rank_loss(prediction, ground_truth, batch_size)
    
    # 组合损失
    loss = reg_loss + alpha * rank_loss
    
    return loss, reg_loss, rank_loss


def set_seed(seed: int = 1234):
    """固定随机种子，便于复现。"""
    random.seed(seed)
    np.random.seed(seed)
    torch.manual_seed(seed)
    if torch.cuda.is_available():
        torch.cuda.manual_seed_all(seed)


# ===================== 数据加载函数 =====================

def load_ticker_list(market_name: str, data_path: str) -> List[str]:
    """从 ticker.csv 或 industry_data.json 加载股票代码列表。"""
    base = os.path.join(data_path, market_name)
    ticker_path = os.path.join(base, f"{market_name.lower()}_ticker.csv")
    if os.path.exists(ticker_path):
        df = pd.read_csv(ticker_path)
        col = df.columns[0]
        ticker_list = df[col].astype(str).str.strip().tolist()
        print(f"✅ 从 {ticker_path} 加载股票列表: {len(ticker_list)} 只")
        return ticker_list
    industry_path = os.path.join(base, f"{market_name.lower()}_industry_data.json")
    if os.path.exists(industry_path):
        with open(industry_path, "r", encoding="utf-8") as f:
            industry_data = json.load(f)
        ticker_list = list(industry_data.keys())
        print(f"✅ 从行业数据文件提取股票列表: {len(ticker_list)} 只")
        return ticker_list
    print(f"⚠️ 未找到 {market_name} 的股票列表文件")
    return []


def load_price_data(
    market_name: str, data_path: str
) -> Tuple[np.ndarray, np.ndarray, np.ndarray, np.ndarray, List[str]]:
    """
    加载价格/技术特征数据。
    返回: eod_data, price_data, gt_data, mask_data, tickers
    """
    dataset_path = os.path.join(data_path, market_name)

    if market_name.upper() == "SP500":
        npy_path = os.path.join(data_path, "SP500", "SP500.npy")
        if not os.path.exists(npy_path):
            raise FileNotFoundError(f"SP500 数据文件不存在: {npy_path}")
        data = np.load(npy_path)
        print(f"✅ 加载 SP500.npy 形状: {data.shape} (股票数, 天数, 特征数)")
        price_data = data[:, :, -1].astype(np.float32)
        mask_data = np.ones((data.shape[0], data.shape[1]), dtype=np.float32)
        gt_data = np.zeros((data.shape[0], data.shape[1]), dtype=np.float32)
        for i in range(data.shape[0]):
            for row in range(PRED_STEPS, data.shape[1]):
                p_prev = data[i, row - PRED_STEPS, -1]
                p_cur = data[i, row, -1]
                if p_prev > 1e-8:
                    gt_data[i, row] = (p_cur - p_prev) / (p_prev + 1e-8)
        tickers = load_ticker_list(market_name, data_path)
        if len(tickers) != data.shape[0]:
            tickers = tickers[: data.shape[0]]
        return data.astype(np.float32), price_data, gt_data, mask_data, tickers

    for name, ext in [("eod_data", "pkl"), ("mask_data", "pkl"), ("gt_data", "pkl"), ("price_data", "pkl")]:
        p = os.path.join(dataset_path, f"{name}.pkl")
        if not os.path.exists(p):
            raise FileNotFoundError(f"数据文件不存在: {p}")
    with open(os.path.join(dataset_path, "eod_data.pkl"), "rb") as f:
        eod_data = pickle.load(f)
    with open(os.path.join(dataset_path, "mask_data.pkl"), "rb") as f:
        mask_data = pickle.load(f)
    with open(os.path.join(dataset_path, "gt_data.pkl"), "rb") as f:
        gt_data = pickle.load(f)
    with open(os.path.join(dataset_path, "price_data.pkl"), "rb") as f:
        price_data = pickle.load(f)
    tickers = load_ticker_list(market_name, data_path)
    if len(tickers) != eod_data.shape[0]:
        tickers = tickers[: eod_data.shape[0]] if tickers else [str(i) for i in range(eod_data.shape[0])]
    print(f"✅ 加载 {market_name} 价格数据: 股票数={eod_data.shape[0]}, 天数={eod_data.shape[1]}")
    return eod_data, price_data, gt_data, mask_data, tickers


def get_trade_dates(
    market_name: str, data_path: str, num_days: int, tickers: List[str]
) -> List[pd.Timestamp]:
    """获取与真实数据列对应的交易日期列表。"""
    base = os.path.join(data_path, market_name)
    for name in ["trade_dates.npy", "trade_dates.csv"]:
        p = os.path.join(base, name)
        if os.path.exists(p):
            if name.endswith(".npy"):
                arr = np.load(p, allow_pickle=True)
                if arr.size >= num_days:
                    dates = pd.to_datetime(arr.ravel()[:num_days]).tolist()
                    print(f"✅ 从 {p} 加载 {len(dates)} 个交易日期")
                    return dates
            else:
                df = pd.read_csv(p)
                dc = "date" if "date" in df.columns else df.columns[0]
                dates = pd.to_datetime(df[dc].dropna()).tolist()
                if len(dates) >= num_days:
                    print(f"✅ 从 {p} 加载 {len(dates)} 个交易日期")
                    return dates[:num_days]
    bdays = pd.date_range("2020-01-01", "2024-12-31", freq="B")
    if len(bdays) < num_days:
        bdays = pd.date_range("2018-01-01", "2024-12-31", freq="B")
    dates = bdays[:num_days].tolist()
    print(f"✅ 使用默认交易日范围，共 {len(dates)} 天")
    return dates


def load_sentiment_data(market_name: str, data_path: str) -> pd.DataFrame:
    """加载情绪特征数据。"""
    combined_path = os.path.join(data_path, market_name, "news_sentiment", f"{market_name.lower()}_combined_daily_sentiment.csv")

    if os.path.exists(combined_path):
        print(f"✅ 使用合并日度情绪文件: {combined_path}")
        df = pd.read_csv(combined_path)
    else:
        raise RuntimeError(f"未找到情绪数据文件: {combined_path}")

    if "ticker" not in df.columns:
        raise RuntimeError("情绪数据中未找到 'ticker' 列")
    if "date" not in df.columns:
        for col in df.columns:
            if col.lower() in ["dt", "trade_date"]:
                df = df.rename(columns={col: "date"})
                break
        if "date" not in df.columns:
            raise RuntimeError("情绪数据中未找到 'date' 列")

    df["date"] = pd.to_datetime(df["date"])
    return df


# 情绪特征列
SENTIMENT_BASE_COLS = [
    "sentiment_polarity_mean",
    "sentiment_polarity_std",
    "sentiment_polarity_count",
    "sentiment_confidence_mean",
    "urgency_score_mean",
    "keywords_earnings_sum",
    "keywords_acquisition_sum",
    "keywords_regulation_sum",
    "keywords_partnership_sum",
    "keywords_product_sum",
    "keywords_financial_sum",
    "keywords_leadership_sum",
    "keywords_lawsuit_sum",
    "text_length_mean",
]
SENTIMENT_EXTRA_COLS = [
    "market_sentiment_mean",
    "market_sentiment_std",
    "relative_sentiment",
    "market_news_volume",
    "relative_news_volume",
    "sentiment_divergence",
]


def build_samples(
    eod_data: np.ndarray,
    sentiment_df: pd.DataFrame,
    gt_data: np.ndarray,
    mask_data: np.ndarray,
    tickers: List[str],
    trade_dates: List[pd.Timestamp],
    lookback_length: int,
    steps: int,
) -> Tuple[List[Dict], List[str], int]:
    """
    构建融合样本：价格特征 + 情绪特征。
    返回: samples, tickers, num_dates
    """
    sentiment_cols = [c for c in SENTIMENT_BASE_COLS + SENTIMENT_EXTRA_COLS if c in sentiment_df.columns]
    if not sentiment_cols:
        raise RuntimeError("情绪数据中未找到可用的情绪特征列")
    print(f"🧠 使用情绪特征列 {len(sentiment_cols)} 个")

    sentiment_df = sentiment_df.copy()
    sentiment_df["date_norm"] = pd.to_datetime(sentiment_df["date"]).dt.normalize()
    date_to_idx = {pd.Timestamp(d).normalize(): i for i, d in enumerate(trade_dates)}
    stock_num, T = gt_data.shape

    # 构建情绪张量
    senti_tensor = np.zeros((stock_num, T, len(sentiment_cols)), dtype=np.float32)
    for stock_idx, ticker in enumerate(tickers):
        sub = sentiment_df[sentiment_df["ticker"] == ticker]
        if sub.empty:
            continue
        for _, row in sub.iterrows():
            d = row["date_norm"]
            if d in date_to_idx:
                t_idx = date_to_idx[d]
                for j, c in enumerate(sentiment_cols):
                    v = row.get(c, np.nan)
                    try:
                        x = float(v) if (pd.notna(v) and v != "") else 0.0
                    except (TypeError, ValueError):
                        x = 0.0
                    senti_tensor[stock_idx, t_idx, j] = np.clip(x, -1e6, 1e6)

    # 构建样本
    samples: List[Dict] = []
    price_feat_dim = eod_data.shape[2]  # 价格特征维度

    for stock_idx in range(stock_num):
        for t_idx in range(lookback_length + steps, T):
            if mask_data[stock_idx, t_idx] < 0.5:
                continue

            # 价格特征窗口
            price_window = eod_data[stock_idx, t_idx - lookback_length : t_idx, :]  # (L, F_price)
            if np.any(np.isnan(price_window)):
                continue

            # 情绪特征窗口
            sentiment_window = senti_tensor[stock_idx, t_idx - lookback_length : t_idx, :]  # (L, F_sent)
            if np.any(np.isnan(sentiment_window)):
                continue

            target = float(gt_data[stock_idx, t_idx])
            samples.append({
                "price_features": price_window.copy(),
                "sentiment_features": sentiment_window.copy(),
                "target": target,
                "stock_idx": stock_idx,
                "time_idx": t_idx,
            })

    if not samples:
        raise RuntimeError("未能构造任何样本，请检查数据")
    print(f"✅ 共构造样本数: {len(samples)} (股票数={stock_num}, 天数={T})")
    return samples, tickers, T


def compute_train_valid_test_index(
    trade_dates: int,
    lookback_length: int,
    steps: int,
    train_ratio: float = 0.6,
    valid_ratio: float = 0.2,
) -> Tuple[int, int]:
    """计算验证集、测试集起始索引。"""
    min_valid_days = max(50, lookback_length + steps + 10)
    min_test_days = max(50, lookback_length + steps + 10)
    min_train_days = max(100, lookback_length + steps + 20)
    valid_index = max(min_train_days, math.floor(trade_dates * train_ratio))
    test_index = max(valid_index + min_valid_days, math.floor(trade_dates * (train_ratio + valid_ratio)))
    if trade_dates - test_index < min_test_days:
        test_index = trade_dates - min_test_days
        if test_index <= valid_index:
            available_days = trade_dates - min_train_days
            if available_days >= min_valid_days + min_test_days:
                valid_index = min_train_days
                test_index = trade_dates - min_test_days
            else:
                raise ValueError(f"数据集太小({trade_dates}天)")
    if valid_index >= test_index or test_index >= trade_dates:
        raise ValueError(f"分割索引无效")
    return valid_index, test_index


def split_samples_by_time(
    samples: List[Dict],
    valid_index: int,
    test_index: int,
) -> Tuple[List[Dict], List[Dict], List[Dict]]:
    """按时间划分样本。"""
    train_samples, valid_samples, test_samples = [], [], []
    for s in samples:
        t = s["time_idx"]
        if t < valid_index:
            train_samples.append(s)
        elif t < test_index:
            valid_samples.append(s)
        else:
            test_samples.append(s)
    print(f"📆 时间划分: valid_index={valid_index}, test_index={test_index}")
    print(f"   训练样本: {len(train_samples)}, 验证样本: {len(valid_samples)}, 测试样本: {len(test_samples)}")
    return train_samples, valid_samples, test_samples


def compute_feature_norm_params(train_samples: List[Dict]) -> Tuple[np.ndarray, np.ndarray, np.ndarray, np.ndarray]:
    """计算价格和情绪特征的标准化参数。"""
    price_feats = [s["price_features"] for s in train_samples]
    senti_feats = [s["sentiment_features"] for s in train_samples]

    all_price = np.concatenate(price_feats, axis=0)
    all_senti = np.concatenate(senti_feats, axis=0)

    price_mean = all_price.mean(axis=0, keepdims=False)
    price_std = all_price.std(axis=0, keepdims=False)
    price_std[price_std == 0] = 1.0

    senti_mean = all_senti.mean(axis=0, keepdims=False)
    senti_std = all_senti.std(axis=0, keepdims=False)
    senti_std[senti_std == 0] = 1.0

    print("✅ 已根据训练集计算特征标准化参数")
    return price_mean.astype(np.float32), price_std.astype(np.float32), senti_mean.astype(np.float32), senti_std.astype(np.float32)


# ===================== 数据集类 =====================

class ConcatMultimodalDataset(Dataset):
    """简单拼接多模态数据集。"""

    def __init__(
        self,
        samples: List[Dict],
        price_mean: np.ndarray,
        price_std: np.ndarray,
        senti_mean: np.ndarray,
        senti_std: np.ndarray,
    ):
        self.samples = samples
        self.price_mean = price_mean.astype(np.float32)
        self.price_std = price_std.astype(np.float32)
        self.senti_mean = senti_mean.astype(np.float32)
        self.senti_std = senti_std.astype(np.float32)

    def __len__(self) -> int:
        return len(self.samples)

    def __getitem__(self, idx: int):
        sample = self.samples[idx]

        # 价格特征标准化
        price_x = sample["price_features"].astype(np.float32)
        price_x = (price_x - self.price_mean) / self.price_std

        # 情绪特征标准化
        senti_x = sample["sentiment_features"].astype(np.float32)
        senti_x = (senti_x - self.senti_mean) / self.senti_std

        y = np.float32(sample["target"])
        stock_idx = int(sample["stock_idx"])
        time_idx = int(sample["time_idx"])

        return (
            torch.from_numpy(price_x),
            torch.from_numpy(senti_x),
            torch.from_numpy(np.array([y], dtype=np.float32)),
            torch.from_numpy(np.array([stock_idx], dtype=np.int64)),
            torch.from_numpy(np.array([time_idx], dtype=np.int64)),
        )


# ===================== 简单拼接融合模型 =====================

class PriceEncoder(nn.Module):
    """价格/技术特征编码器 (LSTM)。"""

    def __init__(
        self,
        input_dim: int,
        hidden_size: int = 64,
        num_layers: int = 2,
        dropout: float = 0.2,
    ):
        super(PriceEncoder, self).__init__()
        self.lstm = nn.LSTM(
            input_size=input_dim,
            hidden_size=hidden_size,
            num_layers=num_layers,
            dropout=dropout if num_layers > 1 else 0.0,
            batch_first=True,
        )
        self.output_dim = hidden_size

    def forward(self, x: torch.Tensor) -> torch.Tensor:
        """x: (batch, seq_len, input_dim)"""
        lstm_out, _ = self.lstm(x)
        return lstm_out[:, -1, :]


class SentimentEncoder(nn.Module):
    """情绪特征编码器 (LSTM)。"""

    def __init__(
        self,
        input_dim: int,
        hidden_size: int = 32,
        num_layers: int = 2,
        dropout: float = 0.2,
    ):
        super(SentimentEncoder, self).__init__()
        self.lstm = nn.LSTM(
            input_size=input_dim,
            hidden_size=hidden_size,
            num_layers=num_layers,
            dropout=dropout if num_layers > 1 else 0.0,
            batch_first=True,
        )
        self.output_dim = hidden_size

    def forward(self, x: torch.Tensor) -> torch.Tensor:
        """x: (batch, seq_len, input_dim)"""
        lstm_out, _ = self.lstm(x)
        return lstm_out[:, -1, :]


class ConcatFusion(nn.Module):
    """
    简单拼接融合模块。
    将价格特征和情绪特征直接拼接，通过全连接层输出。
    """

    def __init__(
        self,
        price_dim: int,
        sentiment_dim: int,
        concat_hidden: int = 64,
    ):
        super(ConcatFusion, self).__init__()

        fused_dim = price_dim + sentiment_dim
        self.fusion_fc = nn.Sequential(
            nn.Linear(fused_dim, concat_hidden),
            nn.GELU(),
            nn.Dropout(0.2),
            nn.Linear(concat_hidden, concat_hidden // 2),
            nn.GELU(),
            nn.Dropout(0.1),
            nn.Linear(concat_hidden // 2, 1)
        )

    def forward(self, price_feat: torch.Tensor, senti_feat: torch.Tensor) -> torch.Tensor:
        concat = torch.cat([price_feat, senti_feat], dim=-1)
        output = self.fusion_fc(concat)
        return output


class ConcatMultimodalModel(nn.Module):
    """简单拼接多模态融合模型。"""

    def __init__(
        self,
        price_input_dim: int,
        sentiment_input_dim: int,
        price_encoder_hidden: int = 64,
        price_encoder_layers: int = 2,
        price_dropout: float = 0.2,
        sentiment_encoder_hidden: int = 32,
        sentiment_encoder_layers: int = 2,
        sentiment_dropout: float = 0.2,
        concat_hidden: int = 64,
    ):
        super(ConcatMultimodalModel, self).__init__()

        # 价格编码器
        self.price_encoder = PriceEncoder(
            input_dim=price_input_dim,
            hidden_size=price_encoder_hidden,
            num_layers=price_encoder_layers,
            dropout=price_dropout,
        )

        # 情绪编码器
        self.sentiment_encoder = SentimentEncoder(
            input_dim=sentiment_input_dim,
            hidden_size=sentiment_encoder_hidden,
            num_layers=sentiment_encoder_layers,
            dropout=sentiment_dropout,
        )

        # 简单拼接融合模块
        self.concat_fusion = ConcatFusion(
            price_dim=price_encoder_hidden,
            sentiment_dim=sentiment_encoder_hidden,
            concat_hidden=concat_hidden,
        )

    def forward(self, price_x: torch.Tensor, senti_x: torch.Tensor) -> torch.Tensor:
        """
        price_x: (batch, seq_len, price_input_dim)
        senti_x: (batch, seq_len, sentiment_input_dim)
        """
        price_feat = self.price_encoder(price_x)
        senti_feat = self.sentiment_encoder(senti_x)
        output = self.concat_fusion(price_feat, senti_feat)
        return output


# ===================== 评估函数 =====================

def evaluate_on_split(
    model: nn.Module,
    dataloader: DataLoader,
    device: torch.device,
    stock_num: int,
    start_day_index: int,
    num_days_in_split: int,
    split_name: str,
) -> Dict[str, float]:
    """在某个数据集划分上进行评估。"""
    model.eval()
    criterion = nn.MSELoss()

    pred_mat = np.zeros((stock_num, num_days_in_split), dtype=np.float32)
    gt_mat = np.zeros((stock_num, num_days_in_split), dtype=np.float32)
    mask_mat = np.zeros((stock_num, num_days_in_split), dtype=np.float32)

    total_loss = 0.0
    total_count = 0

    with torch.no_grad():
        for batch in dataloader:
            price_x, senti_x, y, stock_idx, time_idx = batch
            price_x = price_x.to(device)
            senti_x = senti_x.to(device)
            y = y.to(device)
            stock_idx = stock_idx.to(device)
            time_idx = time_idx.to(device)

            out = model(price_x, senti_x)
            loss = criterion(out, y)

            batch_size = price_x.size(0)
            total_loss += loss.item() * batch_size
            total_count += batch_size

            preds = out.squeeze(1).cpu().numpy()
            gts = y.squeeze(1).cpu().numpy()
            s_idx = stock_idx.squeeze(1).cpu().numpy()
            t_idx = time_idx.squeeze(1).cpu().numpy()

            for i in range(batch_size):
                si = int(s_idx[i])
                ti = int(t_idx[i])
                col = ti - start_day_index
                if 0 <= si < stock_num and 0 <= col < num_days_in_split:
                    pred_mat[si, col] = preds[i]
                    gt_mat[si, col] = gts[i]
                    mask_mat[si, col] = 1.0

    avg_loss = total_loss / max(total_count, 1)
    perf = evaluate(pred_mat, gt_mat, mask_mat)

    if np.isnan(perf.get("IC", 0)):
        perf["IC"] = 0.0
    if np.isnan(perf.get("RIC", 0)):
        perf["RIC"] = 0.0

    print(f"[{split_name}] loss={avg_loss:.2e}  "
          f"mse={perf['mse']:.2e}, IC={perf['IC']:.2e}, "
          f"RIC={perf['RIC']:.2e}, prec@10={perf['prec_10']:.2e}, SR={perf['sharpe5']:.2e}")
    return perf


# ===================== 主函数 =====================

def main():
    set_seed(RANDOM_SEED)
    device = torch.device("cuda" if torch.cuda.is_available() else "cpu")
    print(f"💻 使用设备: {device}")
    print(f"🧪 市场: {MARKET_NAME}, 简单拼接融合模型")
    print("=" * 60)

    # 1. 加载价格数据
    eod_data, price_data, gt_data, mask_data, tickers = load_price_data(MARKET_NAME, DATA_PATH)
    stock_num, num_dates = gt_data.shape

    # 2. 获取交易日期
    trade_dates_list = get_trade_dates(MARKET_NAME, DATA_PATH, num_dates, tickers)

    # 3. 加载情绪数据
    sentiment_df = load_sentiment_data(MARKET_NAME, DATA_PATH)

    # 4. 构建样本
    samples, tickers, num_dates = build_samples(
        eod_data, sentiment_df, gt_data, mask_data, tickers, trade_dates_list, LOOKBACK_LENGTH, PRED_STEPS
    )

    # 5. 数据划分
    valid_index, test_index = compute_train_valid_test_index(
        num_dates, LOOKBACK_LENGTH, PRED_STEPS, train_ratio=0.6, valid_ratio=0.2
    )
    train_samples, valid_samples, test_samples = split_samples_by_time(
        samples, valid_index, test_index
    )

    # 6. 计算特征标准化参数
    price_mean, price_std, senti_mean, senti_std = compute_feature_norm_params(train_samples)

    # 7. 构建数据集
    train_dataset = ConcatMultimodalDataset(train_samples, price_mean, price_std, senti_mean, senti_std)
    valid_dataset = ConcatMultimodalDataset(valid_samples, price_mean, price_std, senti_mean, senti_std)
    test_dataset = ConcatMultimodalDataset(test_samples, price_mean, price_std, senti_mean, senti_std)

    train_loader = DataLoader(train_dataset, batch_size=BATCH_SIZE, shuffle=True, drop_last=False)
    valid_loader = DataLoader(valid_dataset, batch_size=BATCH_SIZE, shuffle=False, drop_last=False)
    test_loader = DataLoader(test_dataset, batch_size=BATCH_SIZE, shuffle=False, drop_last=False)

    # 8. 构建模型
    price_input_dim = train_samples[0]["price_features"].shape[1]
    sentiment_input_dim = train_samples[0]["sentiment_features"].shape[1]

    print(f"📊 价格特征维度: {price_input_dim}, 情绪特征维度: {sentiment_input_dim}")

    model = ConcatMultimodalModel(
        price_input_dim=price_input_dim,
        sentiment_input_dim=sentiment_input_dim,
        price_encoder_hidden=PRICE_ENCODER_HIDDEN,
        price_encoder_layers=PRICE_ENCODER_LAYERS,
        price_dropout=PRICE_DROPOUT,
        sentiment_encoder_hidden=SENTIMENT_ENCODER_HIDDEN,
        sentiment_encoder_layers=SENTIMENT_ENCODER_LAYERS,
        sentiment_dropout=SENTIMENT_DROPOUT,
        concat_hidden=CONCAT_HIDDEN,
    ).to(device)

    optimizer = torch.optim.AdamW(model.parameters(), lr=LEARNING_RATE, weight_decay=WEIGHT_DECAY)
    criterion = nn.MSELoss()

    # 学习率调度器
    scheduler = torch.optim.lr_scheduler.CosineAnnealingWarmRestarts(
        optimizer, T_0=SCHEDULER_T_0, T_mult=2, eta_min=1e-6
    )

    # 使用复合指标选择最佳模型：0.5*IC + 0.5*RIC
    best_valid_metric = -float("inf")
    best_valid_perf = None
    best_test_perf = None
    best_epoch = 0
    no_improve_count = 0

    # 9. 训练循环
    for epoch in range(EPOCHS):
        # Warm-up 阶段
        if epoch < WARMUP_EPOCHS:
            warmup_lr = LEARNING_RATE * (epoch + 1) / WARMUP_EPOCHS
            for param_group in optimizer.param_groups:
                param_group['lr'] = warmup_lr

        model.train()
        total_loss = 0.0
        total_count = 0

        for batch in train_loader:
            price_x, senti_x, y, _, _ = batch
            price_x = price_x.to(device)
            senti_x = senti_x.to(device)
            y = y.to(device)

            optimizer.zero_grad()
            out = model(price_x, senti_x)
            # 使用组合损失：MSE + alpha * 排名损失
            loss, reg_loss, rank_loss = compute_combined_loss(out, y, criterion, ALPHA)
            loss.backward()
            torch.nn.utils.clip_grad_norm_(model.parameters(), max_norm=1.0)
            optimizer.step()

            batch_size = price_x.size(0)
            total_loss += loss.item() * batch_size
            total_count += batch_size

        avg_train_loss = total_loss / max(total_count, 1)
        current_lr = optimizer.param_groups[0]['lr']
        print(f"Epoch {epoch+1}/{EPOCHS} - Train loss: {avg_train_loss:.4e} - LR: {current_lr:.2e}")

        # 验证与测试评估
        valid_perf = evaluate_on_split(
            model, valid_loader, device,
            stock_num=len(tickers),
            start_day_index=valid_index,
            num_days_in_split=test_index - valid_index,
            split_name="Valid",
        )
        test_perf = evaluate_on_split(
            model, test_loader, device,
            stock_num=len(tickers),
            start_day_index=test_index,
            num_days_in_split=num_dates - test_index,
            split_name="Test",
        )

        # 使用复合指标选择最佳模型：0.5*IC + 0.5*RIC
        valid_metric = 0.5 * valid_perf["IC"] + 0.5 * valid_perf["RIC"]
        
        if valid_metric > best_valid_metric:
            best_valid_metric = valid_metric
            best_valid_perf = valid_perf
            best_test_perf = test_perf
            best_epoch = epoch + 1
            no_improve_count = 0
        else:
            no_improve_count += 1

        # Early Stopping
        if no_improve_count >= PATIENCE:
            print(f"\n早停触发！连续 {PATIENCE} 个 epoch 没有提升，停止训练。")
            break

        # 更新学习率调度器
        if epoch >= WARMUP_EPOCHS:
            scheduler.step()

    # 10. 打印最优结果
    print("\n训练结束。")
    if best_valid_perf is not None and best_test_perf is not None:
        print(f"基于验证集复合指标(IC+RIC)最优的轮次: Epoch {best_epoch}")
        print("Best Valid performance:",
              "mse={:.2e}, IC={:.2e}, RIC={:.2e}, prec@10={:.2e}, SR={:.2e}".format(
                  best_valid_perf["mse"], best_valid_perf["IC"], best_valid_perf["RIC"],
                  best_valid_perf["prec_10"], best_valid_perf["sharpe5"]
              ))
        print("Corresponding Test performance:",
              "mse={:.2e}, IC={:.2e}, RIC={:.2e}, prec@10={:.2e}, SR={:.2e}".format(
                  best_test_perf["mse"], best_test_perf["IC"], best_test_perf["RIC"],
                  best_test_perf["prec_10"], best_test_perf["sharpe5"]
              ))

    # 11. 写入Excel结果
    if best_valid_perf and best_test_perf and best_epoch > 0:
        model_name_for_excel = "ConcatFusion"
        excel_file_path = '../training_results.xlsx'

        header = ["Model", "Dataset", "Best Epoch",
                  "Valid MSE", "Valid IC", "Valid RIC", "Valid Prec@10", "Valid SR",
                  "Test MSE", "Test IC", "Test RIC", "Test Prec@10", "Test SR",
                  "Lookback Length", "Learning Rate", "Weight Decay",
                  "Price Encoder Hidden", "Price Encoder Layers", "Price Dropout",
                  "Sentiment Encoder Hidden", "Sentiment Encoder Layers", "Sentiment Dropout",
                  "Concat Hidden", "Batch Size", "Total Epochs",
                  "Num Stocks", "Num Dates"
                  ]

        data_row = [model_name_for_excel, MARKET_NAME, best_epoch,
                    best_valid_perf.get('mse', float('nan')),
                    best_valid_perf.get('IC', float('nan')),
                    best_valid_perf.get('RIC', float('nan')),
                    best_valid_perf.get('prec_10', float('nan')),
                    best_valid_perf.get('sharpe5', float('nan')),
                    best_test_perf.get('mse', float('nan')),
                    best_test_perf.get('IC', float('nan')),
                    best_test_perf.get('RIC', float('nan')),
                    best_test_perf.get('prec_10', float('nan')),
                    best_test_perf.get('sharpe5', float('nan')),
                    LOOKBACK_LENGTH, LEARNING_RATE, WEIGHT_DECAY,
                    PRICE_ENCODER_HIDDEN, PRICE_ENCODER_LAYERS, PRICE_DROPOUT,
                    SENTIMENT_ENCODER_HIDDEN, SENTIMENT_ENCODER_LAYERS, SENTIMENT_DROPOUT,
                    CONCAT_HIDDEN, BATCH_SIZE, EPOCHS,
                    len(tickers), num_dates
                    ]

        try:
            if not os.path.exists(excel_file_path):
                workbook = openpyxl.Workbook()
                sheet = workbook.active
                sheet.title = "Training Results"
                sheet.append(header)
                for col_idx, column_cells in enumerate(sheet.columns):
                    length = max(len(str(cell.value)) if cell.value is not None else 0 for cell in column_cells)
                    sheet.column_dimensions[get_column_letter(col_idx + 1)].width = length + 2
            else:
                workbook = openpyxl.load_workbook(excel_file_path)
                sheet = workbook.active

            sheet.append(data_row)
            workbook.save(excel_file_path)

            print(f"\n✅ 结果已写入 Excel: {excel_file_path} (模型: {model_name_for_excel})")
        except Exception as e:
            print(f"❌ 写入 Excel 失败: {e}")
    else:
        print(f"⚠️ 无最优结果可写入 Excel")


if __name__ == "__main__":
    main()
