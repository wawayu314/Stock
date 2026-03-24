#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
纯情绪 LSTM 股票预测训练脚本（使用真实数据）

特点：
1. 只使用新闻情绪特征，不使用价格/技术特征（纯情绪模型）
2. 使用与 train.py / load_data.py 相同的真实数据源与划分方式：
   - SP500：../dataset/SP500/SP500.npy + sp500_ticker.csv
   - 其他市场：../dataset/{MARKET}/eod_data.pkl, mask_data.pkl, gt_data.pkl, price_data.pkl + ticker 列表
3. 收益率与掩码来自真实数据（gt_data、mask_data），与论文第一部分一致
4. 训练/验证/测试划分与 train.py 一致（train_ratio=0.6, valid_ratio=0.2，最小天数约束）
5. 评估接口：evaluator.evaluate(pred, gt, mask)
"""

import os
import math
import random
import pickle
import json
from typing import List, Dict, Tuple, Optional

import numpy as np
import pandas as pd
import torch
import torch.nn as nn
from torch.utils.data import Dataset, DataLoader

import openpyxl
from openpyxl.utils import get_column_letter

from evaluator import evaluate  # 评估接口与原 train.py 保持一致


# ===================== 全局配置 =====================

# 与 train.py 保持一致：可选 'NASDAQ' / 'SP500' / 'NYSE' / 'A_SHARE'
MARKET_NAME = "SP500"
LOOKBACK_LENGTH = 30          # LSTM 序列长度（使用过去多少天的情绪）
PRED_STEPS = 1                # 预测步长，与 train.py 的 steps=1 一致
DATA_PATH = "../dataset"      # 与 train.py 的 data_path 一致

BATCH_SIZE = 64
EPOCHS = 50
LEARNING_RATE = 1e-3
WEIGHT_DECAY = 1e-4

HIDDEN_SIZE = 64              # LSTM 隐状态维度
NUM_LAYERS = 2                # LSTM 层数
DROPOUT = 0.2                 # LSTM dropout

RANDOM_SEED = 1234


def set_seed(seed: int = 1234):
    """固定随机种子，便于复现。"""
    random.seed(seed)
    np.random.seed(seed)
    torch.manual_seed(seed)
    if torch.cuda.is_available():
        torch.cuda.manual_seed_all(seed)


class SentimentSequenceDataset(Dataset):
    """情绪时间序列数据集（一个样本 = 单只股票过去 L 天情绪 → 下一天收益率）。"""

    def __init__(
        self,
        samples: List[Dict],
        feature_mean: np.ndarray,
        feature_std: np.ndarray,
    ):
        """
        samples 列表中每个元素格式：
        {
            'features': np.ndarray (seq_len, feat_dim),
            'target': float,
            'stock_idx': int,
            'time_idx': int
        }
        feature_mean, feature_std：基于训练集计算的标准化参数。
        """
        self.samples = samples
        self.feature_mean = feature_mean.astype(np.float32)
        self.feature_std = feature_std.astype(np.float32)

    def __len__(self) -> int:
        return len(self.samples)

    def __getitem__(self, idx: int):
        sample = self.samples[idx]
        x = sample["features"].astype(np.float32)  # (seq_len, feat_dim)
        # 特征标准化（使用训练集均值和方差）
        x = (x - self.feature_mean) / self.feature_std
        y = np.float32(sample["target"])
        stock_idx = int(sample["stock_idx"])
        time_idx = int(sample["time_idx"])
        return (
            torch.from_numpy(x),          # (seq_len, feat_dim)
            torch.from_numpy(np.array([y], dtype=np.float32)),  # (1,)
            torch.from_numpy(np.array([stock_idx], dtype=np.int64)),  # (1,)
            torch.from_numpy(np.array([time_idx], dtype=np.int64)),   # (1,)
        )


class SentimentLSTM(nn.Module):
    """纯情绪 LSTM 回归模型：输入情绪序列，输出未来收益率。"""

    def __init__(
        self,
        input_dim: int,
        hidden_size: int = 64,
        num_layers: int = 2,
        dropout: float = 0.2,
    ):
        super(SentimentLSTM, self).__init__()
        self.lstm = nn.LSTM(
            input_size=input_dim,
            hidden_size=hidden_size,
            num_layers=num_layers,
            dropout=dropout if num_layers > 1 else 0.0,
            batch_first=True,
        )
        self.fc = nn.Linear(hidden_size, 1)

    def forward(self, x: torch.Tensor) -> torch.Tensor:
        """
        x: (batch_size, seq_len, input_dim)
        输出: (batch_size, 1)
        """
        lstm_out, _ = self.lstm(x)           # (batch_size, seq_len, hidden_size)
        last_hidden = lstm_out[:, -1, :]     # 取最后一个时间步
        out = self.fc(last_hidden)           # (batch_size, 1)
        return out


def load_ticker_list(market_name: str, data_path: str) -> List[str]:
    """与 train.py 一致：从 ticker.csv 或 industry_data.json 加载股票代码列表（顺序与真实数据一致）。"""
    base = os.path.join(data_path, market_name)
    ticker_path = os.path.join(base, f"{market_name.lower()}_ticker.csv")
    if os.path.exists(ticker_path):
        df = pd.read_csv(ticker_path)
        col = df.columns[0]
        ticker_list = df[col].astype(str).str.strip().tolist()
        print(f"✅ 从 {ticker_path} 加载股票列表: {len(ticker_list)} 只")
        return ticker_list
    industry_path = os.path.join(base, f"{market_name.lower()}_industry_data.json")
    if os.path.exists(industry_path):
        with open(industry_path, "r", encoding="utf-8") as f:
            industry_data = json.load(f)
        ticker_list = list(industry_data.keys())
        print(f"✅ 从行业数据文件提取股票列表: {len(ticker_list)} 只")
        return ticker_list
    print(f"⚠️ 未找到 {market_name} 的股票列表文件")
    return []


def load_real_market_data(
    market_name: str, data_path: str
) -> Tuple[np.ndarray, np.ndarray, np.ndarray, np.ndarray, List[str]]:
    """
    与 train.py / load_data.py 一致：加载真实行情与收益率数据。
    返回: eod_data, price_data, gt_data, mask_data, tickers
    - SP500: ../dataset/SP500/SP500.npy，ticker 顺序与 sp500_ticker.csv 一致
    - 其他: ../dataset/{MARKET}/eod_data.pkl, mask_data.pkl, gt_data.pkl, price_data.pkl
    """
    dataset_path = os.path.join(data_path, market_name)

    if market_name.upper() == "SP500":
        npy_path = os.path.join(data_path, "SP500", "SP500.npy")
        if not os.path.exists(npy_path):
            raise FileNotFoundError(f"SP500 数据文件不存在: {npy_path}")
        data = np.load(npy_path)
        print(f"✅ 加载 SP500.npy 形状: {data.shape} (股票数, 天数, 特征数)")
        price_data = data[:, :, -1].astype(np.float32)
        mask_data = np.ones((data.shape[0], data.shape[1]), dtype=np.float32)
        gt_data = np.zeros((data.shape[0], data.shape[1]), dtype=np.float32)
        for i in range(data.shape[0]):
            for row in range(PRED_STEPS, data.shape[1]):
                p_prev = data[i, row - PRED_STEPS, -1]
                p_cur = data[i, row, -1]
                if p_prev > 1e-8:
                    gt_data[i, row] = (p_cur - p_prev) / (p_prev + 1e-8)
        tickers = load_ticker_list(market_name, data_path)
        if len(tickers) != data.shape[0]:
            tickers = tickers[: data.shape[0]]
        return data.astype(np.float32), price_data, gt_data, mask_data, tickers

    # NASDAQ / NYSE / A_SHARE: 从 pkl 加载
    for name, ext in [("eod_data", "pkl"), ("mask_data", "pkl"), ("gt_data", "pkl"), ("price_data", "pkl")]:
        p = os.path.join(dataset_path, f"{name}.pkl")
        if not os.path.exists(p):
            raise FileNotFoundError(f"真实数据文件不存在: {p}，请先运行 data_collection 生成。")
    with open(os.path.join(dataset_path, "eod_data.pkl"), "rb") as f:
        eod_data = pickle.load(f)
    with open(os.path.join(dataset_path, "mask_data.pkl"), "rb") as f:
        mask_data = pickle.load(f)
    with open(os.path.join(dataset_path, "gt_data.pkl"), "rb") as f:
        gt_data = pickle.load(f)
    with open(os.path.join(dataset_path, "price_data.pkl"), "rb") as f:
        price_data = pickle.load(f)
    tickers = load_ticker_list(market_name, data_path)
    if len(tickers) != eod_data.shape[0]:
        tickers = tickers[: eod_data.shape[0]] if tickers else [str(i) for i in range(eod_data.shape[0])]
    print(f"✅ 加载 {market_name} 真实数据: 股票数={eod_data.shape[0]}, 天数={eod_data.shape[1]}")
    return eod_data, price_data, gt_data, mask_data, tickers


def get_trade_dates(
    market_name: str, data_path: str, num_days: int, tickers: List[str]
) -> List[pd.Timestamp]:
    """
    获取与真实数据列对应的交易日期列表（用于与情绪 CSV 按日期对齐）。
    尝试顺序：trade_dates.npy/csv -> 市场 csv 目录下首只股票日期 -> 默认 2020-2024 交易日。
    """
    base = os.path.join(data_path, market_name)
    # 1) 显式保存的交易日
    for name in ["trade_dates.npy", "trade_dates.csv"]:
        p = os.path.join(base, name)
        if os.path.exists(p):
            if name.endswith(".npy"):
                arr = np.load(p, allow_pickle=True)
                if arr.size >= num_days:
                    dates = pd.to_datetime(arr.ravel()[:num_days]).tolist()
                    print(f"✅ 从 {p} 加载 {len(dates)} 个交易日期")
                    return dates
            else:
                df = pd.read_csv(p)
                dc = "date" if "date" in df.columns else df.columns[0]
                dates = pd.to_datetime(df[dc].dropna()).tolist()
                if len(dates) >= num_days:
                    print(f"✅ 从 {p} 加载 {len(dates)} 个交易日期")
                    return dates[:num_days]
    # 2) 从 csv 目录下首只股票 CSV 的日期列/索引
    csv_dir = os.path.join(base, "csv")
    if tickers and os.path.isdir(csv_dir):
        for t in tickers:
            csv_path = os.path.join(csv_dir, f"{t}.csv")
            if os.path.exists(csv_path):
                try:
                    df = pd.read_csv(csv_path, nrows=num_days + 500)
                    if "date" in df.columns:
                        df["date"] = pd.to_datetime(df["date"])
                        dates = df["date"].dropna().unique().tolist()
                    else:
                        dates = pd.to_datetime(df.index[: num_days + 500]).tolist()
                    if len(dates) >= num_days:
                        dates = sorted(set(dates))[:num_days]
                        print(f"✅ 从 {csv_path} 推断 {len(dates)} 个交易日期")
                        return dates
                except Exception:
                    pass
                break
    # 3) 默认：2020-01-01 至 2024-12-31 的交易日，取前 num_days 个
    bdays = pd.date_range("2020-01-01", "2024-12-31", freq="B")
    if len(bdays) < num_days:
        bdays = pd.date_range("2018-01-01", "2024-12-31", freq="B")
    dates = bdays[:num_days].tolist()
    print(f"✅ 使用默认交易日范围，共 {len(dates)} 天")
    return dates


def load_sentiment_data(market_name: str, data_path: str) -> pd.DataFrame:
    """
    加载情绪特征数据（与 dataset 下新闻情绪数据一致）。
    优先: {data_path}/processed_sentiment/{market}_sentiment.csv
    否则: {data_path}/{market}/news_sentiment/{market_lower}_combined_daily_sentiment.csv
    """
    base = os.path.join(data_path, "processed_sentiment")
    processed_path = os.path.join(base, f"{market_name.lower()}_sentiment.csv")
    combined_path = os.path.join(data_path, market_name, "news_sentiment", f"{market_name.lower()}_combined_daily_sentiment.csv")

    if os.path.exists(processed_path):
        print(f"✅ 使用处理后的情绪特征文件: {processed_path}")
        df = pd.read_csv(processed_path)
    elif os.path.exists(combined_path):
        print(f"✅ 使用合并日度情绪文件: {combined_path}")
        df = pd.read_csv(combined_path)
    else:
        raise RuntimeError(
            f"未找到情绪数据文件，请确认以下路径至少存在一个：\n"
            f" - {processed_path}\n"
            f" - {combined_path}\n"
        )

    # 统一列名
    if "ticker" not in df.columns:
        raise RuntimeError("情绪数据中未找到 'ticker' 列")
    if "date" not in df.columns:
        # 兼容可能的日期列名
        for col in df.columns:
            if col.lower() in ["dt", "trade_date"]:
                df = df.rename(columns={col: "date"})
                break
        if "date" not in df.columns:
            raise RuntimeError("情绪数据中未找到 'date' 列")

    df["date"] = pd.to_datetime(df["date"])
    return df


# 情绪特征列（与 dataset 下 *_daily_sentiment / *_combined_daily_sentiment 列名一致）
SENTIMENT_BASE_COLS = [
    "sentiment_polarity_mean",
    "sentiment_polarity_std",
    "sentiment_polarity_count",
    "sentiment_confidence_mean",
    "urgency_score_mean",
    "keywords_earnings_sum",
    "keywords_acquisition_sum",
    "keywords_regulation_sum",
    "keywords_partnership_sum",
    "keywords_product_sum",
    "keywords_financial_sum",
    "keywords_leadership_sum",
    "keywords_lawsuit_sum",
    "text_length_mean",
]
SENTIMENT_EXTRA_COLS = [
    "market_sentiment_mean",
    "market_sentiment_std",
    "relative_sentiment",
    "market_news_volume",
    "relative_news_volume",
    "sentiment_divergence",
]


def build_samples_from_real_data(
    sentiment_df: pd.DataFrame,
    gt_data: np.ndarray,
    mask_data: np.ndarray,
    tickers: List[str],
    trade_dates: List[pd.Timestamp],
    lookback_length: int,
    steps: int,
) -> Tuple[List[Dict], List[str], int]:
    """
    基于真实数据的 gt_data / mask_data 与情绪数据构造样本。
    与 train.py 使用相同的 (stock_idx, time_idx) 与收益率定义。
    返回: samples, tickers, num_dates
    """
    sentiment_cols = [c for c in SENTIMENT_BASE_COLS + SENTIMENT_EXTRA_COLS if c in sentiment_df.columns]
    if not sentiment_cols:
        raise RuntimeError("情绪数据中未找到可用的情绪特征列，请检查列名。")
    print(f"🧠 使用情绪特征列 {len(sentiment_cols)} 个")

    # 情绪按 (ticker, date) 建立查找表；日期归一化为 date 便于匹配
    sentiment_df = sentiment_df.copy()
    sentiment_df["date_norm"] = pd.to_datetime(sentiment_df["date"]).dt.normalize()
    date_to_idx = {pd.Timestamp(d).normalize(): i for i, d in enumerate(trade_dates)}
    stock_num, T = gt_data.shape
    if len(trade_dates) != T:
        raise RuntimeError(f"trade_dates 长度 {len(trade_dates)} 与数据天数 {T} 不一致")

    # 为每只股票、每个时间点拉取情绪特征，缺失填 0
    senti_tensor = np.zeros((stock_num, T, len(sentiment_cols)), dtype=np.float32)
    for stock_idx, ticker in enumerate(tickers):
        sub = sentiment_df[sentiment_df["ticker"] == ticker]
        if sub.empty:
            continue
        for _, row in sub.iterrows():
            d = row["date_norm"]
            if d in date_to_idx:
                t_idx = date_to_idx[d]
                for j, c in enumerate(sentiment_cols):
                    v = row.get(c, np.nan)
                    try:
                        x = float(v) if (pd.notna(v) and v != "") else 0.0
                    except (TypeError, ValueError):
                        x = 0.0
                    senti_tensor[stock_idx, t_idx, j] = np.clip(x, -1e6, 1e6)

    samples: List[Dict] = []
    for stock_idx in range(stock_num):
        for t_idx in range(lookback_length + steps, T):
            if mask_data[stock_idx, t_idx] < 0.5:
                continue
            window = senti_tensor[stock_idx, t_idx - lookback_length : t_idx, :]  # (L, F)
            if np.any(np.isnan(window)):
                continue
            target = float(gt_data[stock_idx, t_idx])
            samples.append({
                "features": window.copy(),
                "target": target,
                "stock_idx": stock_idx,
                "time_idx": t_idx,
            })

    if not samples:
        raise RuntimeError("未能构造任何样本，请检查情绪数据与真实数据的日期、ticker 是否重叠。")
    print(f"✅ 共构造样本数: {len(samples)} (股票数={stock_num}, 天数={T})")
    return samples, tickers, T


def compute_train_valid_test_index(
    trade_dates: int,
    lookback_length: int,
    steps: int,
    train_ratio: float = 0.6,
    valid_ratio: float = 0.2,
) -> Tuple[int, int]:
    """
    与 train.py 完全一致：计算验证集、测试集起始索引（0-indexed）。
    返回 valid_index, test_index（训练 [0, valid_index)，验证 [valid_index, test_index)，测试 [test_index, trade_dates)）。
    """
    min_valid_days = max(50, lookback_length + steps + 10)
    min_test_days = max(50, lookback_length + steps + 10)
    min_train_days = max(100, lookback_length + steps + 20)
    valid_index = max(min_train_days, math.floor(trade_dates * train_ratio))
    test_index = max(valid_index + min_valid_days, math.floor(trade_dates * (train_ratio + valid_ratio)))
    if trade_dates - test_index < min_test_days:
        test_index = trade_dates - min_test_days
        if test_index <= valid_index:
            available_days = trade_dates - min_train_days
            if available_days >= min_valid_days + min_test_days:
                valid_index = min_train_days
                test_index = trade_dates - min_test_days
            else:
                raise ValueError(f"数据集太小({trade_dates}天)，无法满足最小划分要求")
    if valid_index >= test_index or test_index >= trade_dates:
        raise ValueError(f"分割索引无效: valid_index={valid_index}, test_index={test_index}, trade_dates={trade_dates}")
    return valid_index, test_index


def split_samples_by_time(
    samples: List[Dict],
    valid_index: int,
    test_index: int,
    num_dates: int,
) -> Tuple[List[Dict], List[Dict], List[Dict]]:
    """按与 train.py 相同的 valid_index / test_index 将样本划分为训练/验证/测试。"""
    train_samples, valid_samples, test_samples = [], [], []
    for s in samples:
        t = s["time_idx"]
        if t < valid_index:
            train_samples.append(s)
        elif t < test_index:
            valid_samples.append(s)
        else:
            test_samples.append(s)
    print(f"📆 与 train.py 一致的时间划分: valid_index={valid_index}, test_index={test_index}, 总天数={num_dates}")
    print(f"   训练样本: {len(train_samples)}, 验证样本: {len(valid_samples)}, 测试样本: {len(test_samples)}")
    return train_samples, valid_samples, test_samples


def compute_feature_norm_params(train_samples: List[Dict]) -> Tuple[np.ndarray, np.ndarray]:
    """基于训练样本计算特征标准化参数（均值与标准差）。"""
    # 将所有训练样本的窗口合并到一起计算统计量
    feats_list = [s["features"] for s in train_samples]
    all_feats = np.concatenate(feats_list, axis=0)  # (N_total, F)

    mean = all_feats.mean(axis=0, keepdims=False)
    std = all_feats.std(axis=0, keepdims=False)
    std[std == 0] = 1.0

    print("✅ 已根据训练集计算特征标准化参数")
    return mean.astype(np.float32), std.astype(np.float32)


def evaluate_on_split(
    model: nn.Module,
    dataloader: DataLoader,
    device: torch.device,
    stock_num: int,
    start_day_index: int,
    num_days_in_split: int,
    split_name: str,
) -> Dict[str, float]:
    """
    在某个数据集划分上进行评估，并返回 evaluator 指标。
    只传入当前划分的日期切片 (start_day_index 起共 num_days_in_split 天)，避免全零列导致 IC/RIC 为 nan。
    """
    model.eval()
    criterion = nn.MSELoss()

    # 只分配当前划分的天数，与 train.py 的 validate(start_index, end_index) 一致
    pred_mat = np.zeros((stock_num, num_days_in_split), dtype=np.float32)
    gt_mat = np.zeros((stock_num, num_days_in_split), dtype=np.float32)
    mask_mat = np.zeros((stock_num, num_days_in_split), dtype=np.float32)

    total_loss = 0.0
    total_count = 0

    with torch.no_grad():
        for batch in dataloader:
            x, y, stock_idx, time_idx = batch
            x = x.to(device)                    # (B, L, F)
            y = y.to(device)                    # (B, 1)
            stock_idx = stock_idx.to(device)    # (B, 1)
            time_idx = time_idx.to(device)      # (B, 1)

            out = model(x)                      # (B, 1)
            loss = criterion(out, y)

            batch_size = x.size(0)
            total_loss += loss.item() * batch_size
            total_count += batch_size

            # 填充评估矩阵：列下标为相对当前划分的偏移 (time_idx - start_day_index)
            preds = out.squeeze(1).cpu().numpy()
            gts = y.squeeze(1).cpu().numpy()
            s_idx = stock_idx.squeeze(1).cpu().numpy()
            t_idx = time_idx.squeeze(1).cpu().numpy()

            for i in range(batch_size):
                si = int(s_idx[i])
                ti = int(t_idx[i])
                col = ti - start_day_index
                if 0 <= si < stock_num and 0 <= col < num_days_in_split:
                    pred_mat[si, col] = preds[i]
                    gt_mat[si, col] = gts[i]
                    mask_mat[si, col] = 1.0

    avg_loss = total_loss / max(total_count, 1)
    perf = evaluate(pred_mat, gt_mat, mask_mat)

    # 避免 evaluator 中某日方差为 0 导致相关系数为 nan 时污染均值
    if np.isnan(perf.get("IC", 0)):
        perf["IC"] = 0.0
    if np.isnan(perf.get("RIC", 0)):
        perf["RIC"] = 0.0

    print(f"[{split_name}] loss={avg_loss:.2e}  "
          f"mse={perf['mse']:.2e}, IC={perf['IC']:.2e}, "
          f"RIC={perf['RIC']:.2e}, prec@10={perf['prec_10']:.2e}, SR={perf['sharpe5']:.2e}")
    return perf


def main():
    set_seed(RANDOM_SEED)
    device = torch.device("cuda" if torch.cuda.is_available() else "cpu")
    print(f"💻 使用设备: {device}")
    print(f"🧪 市场: {MARKET_NAME}, 纯情绪 LSTM 模型（真实数据，与 train.py 一致）")
    print("=" * 60)

    # 1. 加载真实行情与收益率（与 train.py / load_data.py 一致）
    eod_data, price_data, gt_data, mask_data, tickers = load_real_market_data(MARKET_NAME, DATA_PATH)
    stock_num, num_dates = gt_data.shape

    # 2. 获取交易日期列表（用于与情绪 CSV 按日期对齐）
    trade_dates_list = get_trade_dates(MARKET_NAME, DATA_PATH, num_dates, tickers)

    # 3. 加载情绪数据
    sentiment_df = load_sentiment_data(MARKET_NAME, DATA_PATH)

    # 4. 用真实 gt_data / mask_data 与情绪构造样本
    samples, tickers, num_dates = build_samples_from_real_data(
        sentiment_df, gt_data, mask_data, tickers, trade_dates_list, LOOKBACK_LENGTH, PRED_STEPS
    )

    # 5. 与 train.py 一致的训练/验证/测试划分
    valid_index, test_index = compute_train_valid_test_index(
        num_dates, LOOKBACK_LENGTH, PRED_STEPS, train_ratio=0.6, valid_ratio=0.2
    )
    train_samples, valid_samples, test_samples = split_samples_by_time(
        samples, valid_index, test_index, num_dates
    )

    # 4. 计算训练集特征标准化参数
    feat_mean, feat_std = compute_feature_norm_params(train_samples)

    # 5. 构建 Dataset 与 DataLoader
    train_dataset = SentimentSequenceDataset(train_samples, feat_mean, feat_std)
    valid_dataset = SentimentSequenceDataset(valid_samples, feat_mean, feat_std)
    test_dataset = SentimentSequenceDataset(test_samples, feat_mean, feat_std)

    train_loader = DataLoader(train_dataset, batch_size=BATCH_SIZE, shuffle=True, drop_last=False)
    valid_loader = DataLoader(valid_dataset, batch_size=BATCH_SIZE, shuffle=False, drop_last=False)
    test_loader = DataLoader(test_dataset, batch_size=BATCH_SIZE, shuffle=False, drop_last=False)

    # 6. 构建模型
    input_dim = train_samples[0]["features"].shape[1]
    model = SentimentLSTM(
        input_dim=input_dim,
        hidden_size=HIDDEN_SIZE,
        num_layers=NUM_LAYERS,
        dropout=DROPOUT,
    ).to(device)

    optimizer = torch.optim.AdamW(model.parameters(), lr=LEARNING_RATE, weight_decay=WEIGHT_DECAY)
    criterion = nn.MSELoss()

    best_valid_ic = -float("inf")
    best_valid_perf = None
    best_test_perf = None
    best_epoch = 0

    # 7. 训练循环
    for epoch in range(EPOCHS):
        model.train()
        total_loss = 0.0
        total_count = 0

        for batch in train_loader:
            x, y, stock_idx, time_idx = batch
            x = x.to(device)
            y = y.to(device)

            optimizer.zero_grad()
            out = model(x)
            loss = criterion(out, y)
            loss.backward()
            torch.nn.utils.clip_grad_norm_(model.parameters(), max_norm=1.0)
            optimizer.step()

            batch_size = x.size(0)
            total_loss += loss.item() * batch_size
            total_count += batch_size

        avg_train_loss = total_loss / max(total_count, 1)
        print(f"Epoch {epoch+1}/{EPOCHS} - Train loss: {avg_train_loss:.4e}")

        # 验证与测试评估（只传入对应时间段切片，避免全零列导致 IC/RIC=nan）
        valid_perf = evaluate_on_split(
            model, valid_loader, device,
            stock_num=len(tickers),
            start_day_index=valid_index,
            num_days_in_split=test_index - valid_index,
            split_name="Valid",
        )
        test_perf = evaluate_on_split(
            model, test_loader, device,
            stock_num=len(tickers),
            start_day_index=test_index,
            num_days_in_split=num_dates - test_index,
            split_name="Test",
        )

        if valid_perf["IC"] > best_valid_ic:
            best_valid_ic = valid_perf["IC"]
            best_valid_perf = valid_perf
            best_test_perf = test_perf
            best_epoch = epoch + 1

    # 8. 打印最优结果
    print("训练结束。")
    if best_valid_perf is not None and best_test_perf is not None:
        print(f"基于验证集 IC 最优的轮次: Epoch {best_epoch}")
        print("Best Valid performance:",
              "mse={:.2e}, IC={:.2e}, RIC={:.2e}, prec@10={:.2e}, SR={:.2e}".format(
                  best_valid_perf["mse"], best_valid_perf["IC"], best_valid_perf["RIC"],
                  best_valid_perf["prec_10"], best_valid_perf["sharpe5"]
              ))
        print("Corresponding Test performance:",
              "mse={:.2e}, IC={:.2e}, RIC={:.2e}, prec@10={:.2e}, SR={:.2e}".format(
                  best_test_perf["mse"], best_test_perf["IC"], best_test_perf["RIC"],
                  best_test_perf["prec_10"], best_test_perf["sharpe5"]
              ))

    # --- 开始写入Excel的逻辑 (与 train.py 一致) ---
    if best_valid_perf and best_test_perf and best_epoch > 0:
        model_name_for_excel = "SentimentLSTM"
        excel_file_path = '../training_results.xlsx'

        header = ["Model", "Dataset", "Best Epoch",
                  "Valid MSE", "Valid IC", "Valid RIC", "Valid Prec@10", "Valid SR",
                  "Test MSE", "Test IC", "Test RIC", "Test Prec@10", "Test SR",
                  "Lookback Length", "Learning Rate", "Weight Decay", "Hidden Size",
                  "Num Layers", "Dropout", "Batch Size", "Total Epochs",
                  "Num Stocks", "Num Dates"
                  ]

        data_row = [model_name_for_excel, MARKET_NAME, best_epoch,
                    best_valid_perf.get('mse', float('nan')),
                    best_valid_perf.get('IC', float('nan')),
                    best_valid_perf.get('RIC', float('nan')),
                    best_valid_perf.get('prec_10', float('nan')),
                    best_valid_perf.get('sharpe5', float('nan')),
                    best_test_perf.get('mse', float('nan')),
                    best_test_perf.get('IC', float('nan')),
                    best_test_perf.get('RIC', float('nan')),
                    best_test_perf.get('prec_10', float('nan')),
                    best_test_perf.get('sharpe5', float('nan')),
                    LOOKBACK_LENGTH, LEARNING_RATE, WEIGHT_DECAY, HIDDEN_SIZE,
                    NUM_LAYERS, DROPOUT, BATCH_SIZE, EPOCHS,
                    len(tickers), num_dates
                    ]

        try:
            if not os.path.exists(excel_file_path):
                workbook = openpyxl.Workbook()
                sheet = workbook.active
                sheet.title = "Training Results"
                sheet.append(header)
                for col_idx, column_cells in enumerate(sheet.columns):
                    length = max(len(str(cell.value)) if cell.value is not None else 0 for cell in column_cells)
                    sheet.column_dimensions[get_column_letter(col_idx + 1)].width = length + 2
            else:
                workbook = openpyxl.load_workbook(excel_file_path)
                sheet = workbook.active

            sheet.append(data_row)
            workbook.save(excel_file_path)

            print(f"\n✅ 结果已写入 Excel: {excel_file_path} (模型: {model_name_for_excel})")
        except Exception as e:
            print(f"❌ 写入 Excel 失败: {e}")
    else:
        print(f"⚠️ 无最优结果可写入 Excel")
    # --- 结束写入Excel的逻辑 ---


if __name__ == "__main__":
    main()

